#!/usr/bin/env python3
"""Text-fit check for the generated deck.

Estimates wrapped height per text paragraph (ceil(text_width / available_width) * line_height),
because a single-line estimate flags every slide and gets ignored. Also checks that tables fit
inside their slide, since a table frame is not a text frame and is easy to skip.

    python3 scripts/check_deck_fit.py
"""
from __future__ import annotations

import pathlib
import sys

from pptx import Presentation
from pptx.util import Emu

DECK = pathlib.Path(__file__).resolve().parents[1] / "deck" / "ShramShield-HACKDAY-1.0.pptx"
LINE_SPACING = 1.28
AVG_CHAR_WIDTH_RATIO = 0.50  # Helvetica-ish average glyph width as a fraction of font size
EMU_PER_INCH = 914400


def inches(emu: int) -> float:
    return emu / EMU_PER_INCH


def text_height_in(text_frame, width_in: float) -> float:
    total = 0.0
    for para in text_frame.paragraphs:
        runs = [r for r in para.runs if r.text]
        if not runs:
            total += 0.12
            continue
        size_pt = runs[0].font.size.pt if runs[0].font.size else 18.0
        text = "".join(run.text for run in runs)
        char_width_in = (size_pt * AVG_CHAR_WIDTH_RATIO) / 72
        chars_per_line = max(1, int(width_in / char_width_in))
        lines = max(1, -(-len(text) // chars_per_line))
        total += lines * (size_pt * LINE_SPACING) / 72
        total += (para.space_after.pt / 72) if para.space_after else 0
    return total


def main() -> None:
    prs = Presentation(DECK)
    problems: list[str] = []
    for index, slide in enumerate(prs.slides, 1):
        slide_h = inches(prs.slide_height)
        slide_w = inches(prs.slide_width)
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                box_w = inches(shape.width) - 0.2
                box_h = inches(shape.height)
                need = text_height_in(shape.text_frame, box_w)
                if need > box_h + 0.05:
                    problems.append(
                        f"slide {index}: text overflows its box by {need - box_h:.2f}in "
                        f"({shape.text_frame.text.strip()[:48]!r})"
                    )
                if inches(shape.left) + inches(shape.width) > slide_w - 0.1:
                    problems.append(f"slide {index}: text box runs off the right edge")
                if inches(shape.top) + inches(shape.height) > slide_h - 0.05:
                    problems.append(f"slide {index}: text box runs off the bottom edge")
            elif shape.has_table:
                table_h = sum(inches(row.height) for row in shape.table.rows)
                if inches(shape.top) + table_h > slide_h - 0.1:
                    problems.append(
                        f"slide {index}: table extends to {inches(shape.top) + table_h:.2f}in "
                        f"of {slide_h:.2f}in"
                    )
                for column_index, column in enumerate(shape.table.columns):
                    col_w = inches(column.width) - 0.2
                    for row in shape.table.rows:
                        cell = row.cells[column_index]
                        if not cell.text:
                            continue
                        size_pt = cell.text_frame.paragraphs[0].runs[0].font.size
                        size = size_pt.pt if size_pt else 14.0
                        chars_per_line = max(1, int(col_w / ((size * AVG_CHAR_WIDTH_RATIO) / 72)))
                        if len(cell.text) > chars_per_line * 2:
                            problems.append(
                                f"slide {index}: table cell needs more than 2 lines "
                                f"({cell.text[:44]!r} in a {col_w:.1f}in column)"
                            )

    if problems:
        print(f"{len(problems)} layout problem(s):")
        for problem in problems:
            print(" -", problem)
        sys.exit(1)
    print(f"deck fit check: OK ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
